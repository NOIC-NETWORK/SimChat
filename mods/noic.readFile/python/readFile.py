import os
import json
import time
import datetime
import hashlib
import mimetypes
from typing import Any, Dict, List, Optional, BinaryIO
from simchat_mod import BaseSkill, ToolSpec, hook



UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "..", "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)

SUPPORTED_EXTENSIONS = {
    "text": [".txt", ".md", ".log", ".csv", ".json", ".xml", ".yaml", ".yml", ".ini", ".cfg", ".conf", ".env"],
    "code": [".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".c", ".cpp", ".h", ".hpp", ".cs", ".go", ".rs", ".rb", ".php", ".swift", ".kt", ".scala"],
    "web": [".html", ".htm", ".css", ".scss", ".less", ".vue", ".svelte", ".jsx", ".tsx"],
    "document": [".pdf", ".doc", ".docx", ".rtf", ".odt", ".ods", ".odp"],
    "spreadsheet": [".xls", ".xlsx", ".csv", ".tsv"],
    "presentation": [".ppt", ".pptx", ".key"],
    "image": [".png", ".jpg", ".jpeg", ".gif", ".bmp", ".svg", ".webp", ".ico"],
    "audio": [".mp3", ".wav", ".ogg", ".flac", ".aac", ".m4a"],
    "video": [".mp4", ".avi", ".mov", ".mkv", ".webm", ".flv"],
    "archive": [".zip", ".rar", ".7z", ".tar", ".gz", ".bz2"],
    "notebook": [".ipynb"]
}

MAX_FILE_SIZE = 50 * 1024 * 1024
MAX_TEXT_LENGTH = 100000

try:
    import PyPDF2
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False
    print("[ReadFile] PyPDF2 未安装，PDF 读取功能将受限")

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    print("[ReadFile] python-docx 未安装，Word 文档读取功能将受限")

try:
    import openpyxl
    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False
    print("[ReadFile] openpyxl 未安装，Excel 表格读取功能将受限")


def get_file_category(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    for category, extensions in SUPPORTED_EXTENSIONS.items():
        if ext in extensions:
            return category
    return "unknown"


def is_supported_file(file_path: str) -> bool:
    ext = os.path.splitext(file_path)[1].lower()
    for extensions in SUPPORTED_EXTENSIONS.values():
        if ext in extensions:
            return True
    return False


def get_file_info(file_path: str) -> Dict[str, Any]:
    if not os.path.exists(file_path):
        return {"exists": False, "error": "文件不存在"}

    file_size = os.path.getsize(file_path)
    file_category = get_file_category(file_path)
    file_ext = os.path.splitext(file_path)[1].lower()
    file_hash = ""

    if file_size <= MAX_FILE_SIZE:
        try:
            with open(file_path, "rb") as f:
                file_hash = hashlib.md5(f.read()).hexdigest()
        except Exception:
            pass

    return {
        "exists": True,
        "name": os.path.basename(file_path),
        "path": file_path,
        "size": file_size,
        "size_formatted": format_file_size(file_size),
        "extension": file_ext,
        "category": file_category,
        "supported": is_supported_file(file_path),
        "md5": file_hash if file_size <= MAX_FILE_SIZE else "file_too_large",
        "modified": datetime.datetime.fromtimestamp(os.path.getmtime(file_path)).strftime("%Y-%m-%d %H:%M:%S"),
        "mime_type": mimetypes.guess_type(file_path)[0] or "application/octet-stream"
    }


def format_file_size(size_bytes: int) -> str:
    if size_bytes < 1024:
        return f"{size_bytes} B"
    elif size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    elif size_bytes < 1024 * 1024 * 1024:
        return f"{size_bytes / (1024 * 1024):.1f} MB"
    else:
        return f"{size_bytes / (1024 * 1024 * 1024):.1f} GB"


class TextFileReader:
    @staticmethod
    def read(file_path: str, encoding: str = "utf-8", max_length: int = MAX_TEXT_LENGTH) -> Dict[str, Any]:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                content = f.read(max_length)

            total_size = os.path.getsize(file_path)
            truncated = len(content) >= max_length

            return {
                "success": True,
                "content": content,
                "length": len(content),
                "total_size": total_size,
                "truncated": truncated,
                "encoding": encoding
            }
        except UnicodeDecodeError:
            return TextFileReader.read(file_path, encoding="gbk", max_length=max_length)
        except Exception as e:
            return {"success": False, "error": str(e)}


class PDFReader:
    @staticmethod
    def read(file_path: str, max_pages: int = 10) -> Dict[str, Any]:
        if not PDF_AVAILABLE:
            return {"success": False, "error": "PyPDF2 未安装"}

        try:
            reader = PyPDF2.PdfReader(file_path)
            pages = []
            total_pages = len(reader.pages)
            pages_to_read = min(max_pages, total_pages)

            for i in range(pages_to_read):
                try:
                    page = reader.pages[i]
                    text = page.extract_text() or ""
                    pages.append({"page": i + 1, "content": text[:MAX_TEXT_LENGTH]})
                except Exception as e:
                    pages.append({"page": i + 1, "content": "", "error": str(e)})

            full_text = "\n\n".join([f"--- 第{p['page']}页 ---\n{p['content']}" for p in pages])

            return {
                "success": True,
                "total_pages": total_pages,
                "pages_read": pages_to_read,
                "content": full_text,
                "pages": pages
            }
        except Exception as e:
            return {"success": False, "error": str(e)}


class DocxReader:
    @staticmethod
    def read(file_path: str) -> Dict[str, Any]:
        if not DOCX_AVAILABLE:
            return {"success": False, "error": "python-docx 未安装"}

        try:
            doc = docx.Document(file_path)
            paragraphs = []
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)

            tables = []
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text for cell in row.cells]
                    table_data.append(row_data)
                tables.append(table_data)

            full_text = "\n".join(paragraphs)
            if tables:
                full_text += "\n\n[表格数据]\n"
                for i, table in enumerate(tables):
                    full_text += f"\n表格 {i + 1}:\n"
                    for row in table:
                        full_text += " | ".join(row) + "\n"

            return {
                "success": True,
                "content": full_text,
                "paragraphs_count": len(paragraphs),
                "tables_count": len(tables)
            }
        except Exception as e:
            return {"success": False, "error": str(e)}


class XlsxReader:
    @staticmethod
    def read(file_path: str, sheet_name: Optional[str] = None, max_rows: int = 100) -> Dict[str, Any]:
        if not XLSX_AVAILABLE:
            return {"success": False, "error": "openpyxl 未安装"}

        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            results = []

            sheets_to_read = [sheet_name] if sheet_name else wb.sheetnames[:3]

            for sn in sheets_to_read:
                if sn not in wb.sheetnames:
                    continue

                ws = wb[sn]
                rows = []
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= max_rows:
                        break
                    rows.append([str(cell) if cell is not None else "" for cell in row])

                results.append({
                    "sheet": sn,
                    "rows_count": len(rows),
                    "data": rows
                })

            wb.close()

            return {
                "success": True,
                "sheets_read": len(results),
                "data": results
            }
        except Exception as e:
            return {"success": False, "error": str(e)}


class ImageAnalyzer:
    @staticmethod
    def get_metadata(file_path: str) -> Dict[str, Any]:
        try:
            file_size = os.path.getsize(file_path)
            with open(file_path, "rb") as f:
                header = f.read(32)

            metadata = {
                "success": True,
                "file_size": file_size,
                "file_type": get_file_category(file_path),
                "mime_type": mimetypes.guess_type(file_path)[0]
            }

            if file_path.lower().endswith(".png"):
                metadata["format"] = "PNG"
            elif file_path.lower().endswith((".jpg", ".jpeg")):
                metadata["format"] = "JPEG"
            elif file_path.lower().endswith(".gif"):
                metadata["format"] = "GIF"
            elif file_path.lower().endswith(".svg"):
                metadata["format"] = "SVG"
                metadata["note"] = "SVG 是矢量图，可直接读取文本内容"
                try:
                    with open(file_path, "r", encoding="utf-8") as f:
                        metadata["svg_preview"] = f.read(500)
                except Exception:
                    pass
            else:
                metadata["format"] = "未知图片格式"

            return metadata
        except Exception as e:
            return {"success": False, "error": str(e)}


class NotebookReader:
    @staticmethod
    def read(file_path: str) -> Dict[str, Any]:
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                notebook = json.load(f)

            cells = []
            for i, cell in enumerate(notebook.get("cells", [])):
                cell_type = cell.get("cell_type", "unknown")
                source = "".join(cell.get("source", []))
                cells.append({
                    "index": i,
                    "type": cell_type,
                    "content": source[:MAX_TEXT_LENGTH]
                })

            return {
                "success": True,
                "cells_count": len(cells),
                "cells": cells,
                "metadata": notebook.get("metadata", {})
            }
        except Exception as e:
            return {"success": False, "error": str(e)}


class ReadFileSkill(BaseSkill):
    skill_name = "readFile"
    skill_description = "读取用户上传的文件内容并提供给AI分析，支持多种文件格式"
    skill_version = "1.0.0"
    trigger_keywords = ["读取文件", "查看文件", "分析文件", "上传文件", "file", "read", "读取", "文件内容", "查看", "分析"]
    input_schema = {
        "file_path": str,
        "max_length": int,
        "encoding": str,
        "file_url": str
    }

    def execute(self, params: Dict[str, Any]) -> Dict[str, Any]:
        file_path = params.get("file_path", "")
        max_length = params.get("max_length", MAX_TEXT_LENGTH)
        encoding = params.get("encoding", "utf-8")

        if not file_path:
            return self._handle_file_upload(params)

        if file_path.startswith("http://") or file_path.startswith("https://"):
            return self._download_and_read(file_path, max_length)

        if not os.path.exists(file_path):
            return {"success": False, "error": f"文件不存在: {file_path}"}

        file_size = os.path.getsize(file_path)
        if file_size > MAX_FILE_SIZE:
            return {"success": False, "error": f"文件过大 ({format_file_size(file_size)})，最大支持 {format_file_size(MAX_FILE_SIZE)}"}

        file_category = get_file_category(file_path)
        file_info = get_file_info(file_path)

        readers = {
            "text": lambda: TextFileReader.read(file_path, encoding, max_length),
            "code": lambda: TextFileReader.read(file_path, encoding, max_length),
            "web": lambda: TextFileReader.read(file_path, encoding, max_length),
            "document": lambda: self._read_document(file_path, max_length),
            "spreadsheet": lambda: self._read_spreadsheet(file_path),
            "presentation": lambda: {"success": False, "error": "PPT 文件暂不支持读取"},
            "image": lambda: ImageAnalyzer.get_metadata(file_path),
            "audio": lambda: {"success": False, "error": "音频文件暂不支持直接读取，请先转换为文本"},
            "video": lambda: {"success": False, "error": "视频文件暂不支持直接读取"},
            "archive": lambda: {"success": False, "error": "压缩包请先解压后再读取内部文件"},
            "notebook": lambda: NotebookReader.read(file_path)
        }

        reader_func = readers.get(file_category)
        if reader_func:
            try:
                content_result = reader_func()
                return {
                    "success": content_result.get("success", False),
                    "file_info": file_info,
                    "content": content_result.get("content", content_result),
                    "category": file_category
                }
            except Exception as e:
                return {"success": False, "error": str(e), "file_info": file_info}
        else:
            return TextFileReader.read(file_path, encoding, max_length)

    def _read_document(self, file_path: str, max_length: int) -> Dict[str, Any]:
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".pdf":
            return PDFReader.read(file_path)
        elif ext in [".docx", ".doc"]:
            return DocxReader.read(file_path)
        elif ext == ".txt" or ext == ".md":
            return TextFileReader.read(file_path)
        else:
            return TextFileReader.read(file_path)

    def _read_spreadsheet(self, file_path: str) -> Dict[str, Any]:
        ext = os.path.splitext(file_path)[1].lower()

        if ext in [".xlsx", ".xls"]:
            return XlsxReader.read(file_path)
        elif ext == ".csv":
            return TextFileReader.read(file_path)
        else:
            return TextFileReader.read(file_path)

    def _download_and_read(self, url: str, max_length: int) -> Dict[str, Any]:
        try:
            import urllib.request
            import tempfile

            response = urllib.request.urlopen(url, timeout=30)
            content_type = response.headers.get("Content-Type", "")

            file_name = url.split("/")[-1].split("?")[0] or "downloaded_file"
            temp_path = os.path.join(UPLOAD_DIR, f"temp_{int(time.time())}_{file_name}")

            with open(temp_path, "wb") as f:
                f.write(response.read(MAX_FILE_SIZE))

            result = self.execute({"file_path": temp_path, "max_length": max_length})
            result["downloaded_from"] = url

            try:
                os.remove(temp_path)
            except Exception:
                pass

            return result
        except Exception as e:
            return {"success": False, "error": f"下载失败: {str(e)}"}

    def _handle_file_upload(self, params: Dict[str, Any]) -> Dict[str, Any]:
        file_name = params.get("file_name", "")
        file_content = params.get("file_content", "")
        file_data = params.get("file_data", "")

        if not file_name:
            return {"success": False, "error": "缺少 file_name 参数"}

        if not file_content and not file_data:
            return {"success": False, "error": "缺少文件内容"}

        try:
            file_path = os.path.join(UPLOAD_DIR, f"{int(time.time())}_{file_name}")

            if file_data:
                import base64
                decoded = base64.b64decode(file_data)
                with open(file_path, "wb") as f:
                    f.write(decoded)
            elif file_content:
                with open(file_path, "w", encoding="utf-8") as f:
                    f.write(file_content)

            result = self.execute({"file_path": file_path})
            result["upload_path"] = file_path
            return result
        except Exception as e:
            return {"success": False, "error": f"文件保存失败: {str(e)}"}

    def list_uploaded_files(self) -> Dict[str, Any]:
        files = []
        if os.path.exists(UPLOAD_DIR):
            for fname in os.listdir(UPLOAD_DIR):
                fpath = os.path.join(UPLOAD_DIR, fname)
                if os.path.isfile(fpath):
                    files.append(get_file_info(fpath))
        return {"success": True, "files": files, "count": len(files)}

    def delete_uploaded_file(self, file_name: str) -> Dict[str, Any]:
        file_path = os.path.join(UPLOAD_DIR, file_name)
        if os.path.exists(file_path):
            try:
                os.remove(file_path)
                return {"success": True, "message": f"文件 {file_name} 已删除"}
            except Exception as e:
                return {"success": False, "error": str(e)}
        return {"success": False, "error": "文件不存在"}


def init_readFile_skill():
    manager = get_skill_manager()
    registry = manager.get_registry()

    existing = registry.get_skill("readFile")
    if existing:
        print("[ReadFile] readFile 技能已注册")
        return existing

    skill = ReadFileSkill()
    registry.register_skill(skill)
    print("[ReadFile] readFile 技能初始化完成")
    return skill


if __name__ == "__main__":
    print("[ReadFile] 初始化 readFile 技能...")
    skill = init_readFile_skill()

    print("\n[ReadFile] 支持的文件类型:")
    for category, extensions in SUPPORTED_EXTENSIONS.items():
        print(f"  {category}: {', '.join(extensions)}")

    print("\n[ReadFile] 示例使用:")
    print('  manager = get_skill_manager()')
    print('  result = manager.execute_skill("readFile", {"file_path": "/path/to/file.txt"})')

# ----- 工具实例 -----
_instance = ReadFileSkill()


def execute(params):
    return _instance.execute(params)


@hook("agent.tool")
def register(api):
    api.register_tool(ToolSpec(
        name="readFile",
        label="readFile",
        description=_instance.skill_description,
        risk="read",
        usage="readFile PARAMS: ...",
        input_schema=_instance.input_schema,
        triggers=_instance.trigger_keywords,
        backend="readFile:execute",
    ))